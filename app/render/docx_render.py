"""Word (.docx) rendering of the finalised document HTML.

Offer letters and service agreements get the masthead banner in the running
header. NCNDAs render plain (no letterhead), matching the sample NCNDAs.
"""
import io

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..documents.registry import get_entity
from ..documents import letterhead as lh

ACCENT = {"fouress": RGBColor(0xE0, 0x8B, 0x1E), "estel": RGBColor(0x1F, 0x7A, 0x6D)}


def _accent_for(ent):
    return ACCENT.get(ent.get("brand", "fouress"), ACCENT["fouress"])


def _add_inline(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        text = str(node)
        if text.strip() or text == " ":
            run = paragraph.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    b = bold or node.name in ("strong", "b")
    i = italic or node.name in ("em", "i")
    for child in node.children:
        _add_inline(paragraph, child, b, i)


def _render_block(doc, el, accent):
    name = el.name
    if name in ("h2", "h3"):
        p = doc.add_paragraph()
        _add_inline(p, el)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(13 if name == "h2" else 11)
            run.font.color.rgb = accent
        if "doc-title" in (el.get("class") or []):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif name == "p":
        p = doc.add_paragraph()
        if "subtitle" in (el.get("class") or []):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_inline(p, el)
    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        for li in el.find_all("li", recursive=False):
            p = doc.add_paragraph(style=style)
            _add_inline(p, li)
    elif name == "table":
        _render_table(doc, el)


def _render_table(doc, el):
    rows = el.find_all("tr", recursive=True)
    if not rows:
        return
    ncols = max(len(r.find_all(["td", "th"], recursive=False)) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for r in rows:
        cells = r.find_all(["td", "th"], recursive=False)
        row = table.add_row()
        for idx, c in enumerate(cells):
            if idx >= ncols:
                break
            cell = row.cells[idx]
            cell.paragraphs[0].text = ""
            first = True
            blocks = c.find_all(["p", "h2", "h3"], recursive=False) or [c]
            for block in blocks:
                para = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                _add_inline(para, block)


def _walk(doc, node, accent):
    for child in node.children:
        if not isinstance(child, Tag):
            continue
        classes = child.get("class") or []
        if child.name == "section" and "annex" in classes:
            doc.add_page_break()
            _walk(doc, child, accent)
        elif child.name in ("h2", "h3", "p", "ul", "ol", "table"):
            _render_block(doc, child, accent)
        else:
            _walk(doc, child, accent)


def _setup_letterhead(doc, ent, accent, letterhead=True):
    sec = doc.sections[0]
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    if not letterhead:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        return
    sec.top_margin = Inches(1.75)
    sec.bottom_margin = Inches(0.9)

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run()
    try:
        run.add_picture(str(lh.masthead_path(ent)), width=Inches(6.4))
    except Exception:
        r = hp.add_run(ent["legal_name"]); r.bold = True; r.font.color.rgb = accent

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(ent["legal_name"] + "  |  " + "  |  ".join(
        ent["address_lines"] + [ent["contact_line"]]))
    fr.font.size = Pt(7)
    fr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)


def render_docx(entity_key, body_html, letterhead=True):
    ent = get_entity(entity_key)
    accent = _accent_for(ent)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    _setup_letterhead(doc, ent, accent, letterhead)

    soup = BeautifulSoup(body_html, "html.parser")
    root = soup.find(id="document") or soup
    _walk(doc, root, accent)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
